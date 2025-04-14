import os
import io
import requests
from flask import Flask, request, send_file, jsonify
from pdf2image import convert_from_path
from google.cloud import vision
from docx import Document
from tempfile import NamedTemporaryFile

app = Flask(__name__)

os.environ["GOOGLE_APPLICATION_CREDENTIALS"] = "google_credentials.json"

# Initialize Google Vision API client
client = vision.ImageAnnotatorClient()

def download_pdf(url):
    response = requests.get(url)
    if response.status_code == 200:
        tmp = NamedTemporaryFile(delete=False, suffix=".pdf")
        tmp.write(response.content)
        tmp.close()
        return tmp.name
    else:
        raise Exception("Failed to download PDF")


def pdf_to_images(pdf_path):
    images = convert_from_path(pdf_path, poppler_path="./poppler/bin")

    return images


def images_to_text(images):
    text_list = []
    for img in images:
        img_byte_arr = io.BytesIO()
        img.save(img_byte_arr, format="PNG")
        content = img_byte_arr.getvalue()

        image = vision.Image(content=content)
        response = client.document_text_detection(image=image)

        if response.error.message:
            raise Exception(f"Vision API error: {response.error.message}")
        text = response.full_text_annotation.text
        text_list.append(text)
    return text_list


def create_docx(text_list):
    doc = Document()
    for idx, text in enumerate(text_list):
        doc.add_heading(f"Page {idx + 1}", level=2)
        doc.add_paragraph(text)
        doc.add_page_break()

    output = NamedTemporaryFile(delete=False, suffix=".docx")
    doc.save(output.name)
    return output.name


@app.route("/process", methods=["POST"])
def process_pdf():
    try:
        data = request.get_json()
        pdf_url = data.get("pdf_url")
        if not pdf_url:
            return jsonify({"error": "Missing 'pdf_url'"}), 400

        pdf_path = download_pdf(pdf_url)
        images = pdf_to_images(pdf_path)
        text_list = images_to_text(images)
        docx_path = create_docx(text_list)

        return send_file(docx_path, as_attachment=True, download_name="output.docx")

    except Exception as e:
        return jsonify({"error": str(e)}), 500

if __name__ == "__main__":
    app.run(host='0.0.0.0', port=5000, debug=True)